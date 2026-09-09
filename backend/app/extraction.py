"""Plain-text extraction from uploaded documents.

Shared by the stateless `/uploads/extract` utility and the chat upload stream,
so the parsing rules, size limits and error messages live in one place. Nothing
here touches the database or the request: give it bytes, get text back.
"""
from __future__ import annotations

import io

from fastapi import HTTPException

# Guardrails. The text cap keeps a 200-page PDF from becoming an unbounded
# message; it sits under the message content limit so a caller still has room
# for a question and the framing around the excerpt.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_TEXT_CHARS = 50_000

_PDF = {"application/pdf"}
_DOCX = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
}


def detect_kind(filename: str | None, content_type: str | None) -> str:
    """Return "pdf" or "docx", or raise 415 for anything else."""
    name = (filename or "").lower()
    if content_type in _PDF or name.endswith(".pdf"):
        return "pdf"
    if content_type in _DOCX or name.endswith(".docx"):
        return "docx"
    raise HTTPException(415, "Only PDF and DOCX files are supported.")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - surface any parse failure as 422
        raise HTTPException(422, "That PDF could not be read; it may be corrupt "
                                 "or password-protected.") from exc
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, "That DOCX could not be read; it may be corrupt "
                                 "or not a real Word document.") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Pull table cell text too; requirements often live in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract(filename: str | None, content_type: str | None,
            data: bytes) -> tuple[str, str, bool]:
    """Validate and extract. Returns (kind, text, truncated).

    Raises HTTPException with a caller-friendly message on an unsupported type
    (415), an oversized file (413), or a file with no readable text (422).
    """
    kind = detect_kind(filename, content_type)
    if not data:
        raise HTTPException(422, "The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            413, f"File is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.")

    text = (_extract_pdf(data) if kind == "pdf" else _extract_docx(data)).strip()
    if not text:
        raise HTTPException(
            422,
            "No selectable text was found. If this is a scanned document it "
            "needs OCR, which is not supported here.",
        )

    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS].rsplit(" ", 1)[0] + "…"
    return kind, text, truncated
